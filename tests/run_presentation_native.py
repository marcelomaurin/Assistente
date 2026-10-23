"""Native Windows presentation checks. Requires Lazarus, a desktop and PowerPoint."""
from pathlib import Path
import io, wave, argparse, hashlib, http.server, json, struct, subprocess, threading, zipfile
from pptx import Presentation
from PIL import Image
def chunk(name,data):
    return name+struct.pack('<I',len(data))+data+(b'\0' if len(data)%2 else b'')
def avi(path):
    w=h=64;n=20;size=w*h*3
    avih=struct.pack('<14I',200000,size*5,0,16,n,0,1,size,w,h,0,0,0,0)
    strh=struct.pack('<4s4sIHH8I4h',b'vids',b'DIB ',0,0,0,0,1,5,0,n,size,0xffffffff,0,0,0,w,h)
    strf=struct.pack('<IiiHHIIiiII',40,w,h,1,24,0,size,0,0,0,0)
    hdrl=chunk(b'LIST',b'hdrl'+chunk(b'avih',avih)+chunk(b'LIST',b'strl'+chunk(b'strh',strh)+chunk(b'strf',strf)))
    frames=[];index=[];offset=4
    for i in range(n):
        data=bytes([40,60+i*6,190])*w*h;frames.append(chunk(b'00db',data))
        index.append(struct.pack('<4sIII',b'00db',16,offset,len(data)));offset+=8+len(data)
    body=b'AVI '+hdrl+chunk(b'LIST',b'movi'+b''.join(frames))+chunk(b'idx1',b''.join(index))
    path.write_bytes(chunk(b'RIFF',body))
def run(cmd,timeout=60):
    p=subprocess.run(list(map(str,cmd)),capture_output=True,timeout=timeout)
    assert p.returncode==0,(cmd,p.returncode,p.stdout.decode(errors='replace'),p.stderr.decode(errors='replace'))
def main():
    p=argparse.ArgumentParser();p.add_argument('build',type=Path);p.add_argument('model',type=Path);p.add_argument('--without-office',action='store_true');a=p.parse_args()
    build=a.build.resolve();root=build/'presentation-native';(root/'assets').mkdir(parents=True,exist_ok=True)
    for n,c in [('one','navy'),('two','darkorange')]:Image.new('RGB',(640,360),c).save(root/f'assets/{n}.png')
    avi(root/'assets/movie.avi');prs=Presentation()
    for i in range(3):
        slide=prs.slides.add_slide(prs.slide_layouts[0]);slide.shapes.title.text=f'Slide {i+1}'
    prs.save(root/'assets/slides.pptx')
    script=('[abrir arquivo=assets/one.png]Uma imagem para observar e pausar. '
            '[articulacao=head,y,20,500,suave]A cabeca se move enquanto a boca fica parada. '
            '[abrir arquivo=assets/two.png]Outra imagem. [fechar arquivo]Fechou. '
            '[abrir arquivo=assets/slides.pptx]Primeiro slide. [proximo slide]Segundo slide. '
            '[slide anterior]Primeiro novamente. [ir para slide=3]Terceiro slide. [fechar arquivo]')
    if a.without_office:script=script[:script.index('[abrir arquivo=assets/slides.pptx]')]+'[fechar arquivo]'
    for name,s in [('main',script),('cancel','[abrir arquivo=assets/one.png]Uma frase longa para testar cancelamento e limpeza de todos os recursos visuais.')]:
        source=root/(name+'.txt');source.write_text(s,encoding='utf-8')
        run([build/'pva_tool.exe','pack',source,root,root/(name+'.pva')])
    for mode in ['media','editor','main']:
        report=root/(mode+'-result.txt');report.unlink(missing_ok=True)
        run([build/'presentation_native_smoke.exe',root,mode,a.model.resolve(),('without-office' if a.without_office else 'with-office')],timeout=65)
        assert report.exists(),mode
        body=report.read_text(encoding='utf-8');assert 'FAIL' not in body,body;print(body.strip(),flush=True)
    original='Olá [abrir arquivo=assets/one.png]mundo.\n'
    seen=[]
    class Handler(http.server.BaseHTTPRequestHandler):
        def do_POST(self):
            body=json.loads(self.rfile.read(int(self.headers['Content-Length'])));seen.append(body)
            if self.path.startswith('/audio'):
                assert '[' not in body['input'],body
                stream=io.BytesIO()
                with wave.open(stream,'wb') as wav:
                    wav.setnchannels(1);wav.setsampwidth(2);wav.setframerate(16000);wav.writeframes(b'\0'*(16000*2*3//4))
                result=stream.getvalue()
                self.send_response(200);self.send_header('Content-Type','audio/wav');self.send_header('Content-Length',str(len(result)));self.end_headers();self.wfile.write(result);return
            value=('[sorriso]'+original) if self.path.startswith('/accept') else 'Texto alterado.'
            result=json.dumps({'choices':[{'message':{'content':json.dumps({'script':value},ensure_ascii=False)}}]}).encode()
            self.send_response(200);self.send_header('Content-Type','application/json');self.send_header('Content-Length',str(len(result)));self.end_headers();self.wfile.write(result)
        def log_message(self,*args):pass
    server=http.server.ThreadingHTTPServer(('127.0.0.1',0),Handler)
    thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    try:
        for mode in ['accept','reject']:
            run([build/'presentation_ai_tests.exe',f'http://127.0.0.1:{server.server_port}/{mode}',mode])
        run([build/'presentation_speech_tests.exe',f'http://127.0.0.1:{server.server_port}/audio'])
        assert len(seen)==5
    finally:server.shutdown();thread.join();server.server_close()
    if a.without_office:print('SKIP: PowerPoint activation/navigation requires an interactive Office session')
    print('PASS: real WAV playback, media/gesture queue, pause/resume, cancellation while paused')
    print('PASS: real configured ChatGPT component against local mock; preserved UTF-8/media; rejected rewrite')
if __name__=='__main__':main()
